from pathlib import Path
import csv, json, math, os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_ROW_HEIGHT_RULE

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'manuscript'/'PG_CARE_IEEE_Transactions_Manuscript.docx'
OUT.parent.mkdir(exist_ok=True)
FIG=ROOT/'figures'; RES=ROOT/'results'; REF=ROOT/'references'/'references_2020_2026.csv'

# Load results
import pandas as pd
cm=pd.read_csv(RES/'classification_metrics.csv')
mb=pd.read_csv(RES/'modern_baselines.csv')
pg=pd.read_csv(RES/'pgcare_lgbm_metrics.csv')
ci=pd.read_csv(RES/'bootstrap_ci_lgbm.csv')
pbd=pd.read_csv(RES/'paired_bootstrap_lgbm.csv')
reg=pd.read_csv(RES/'regression_metrics.csv')
ood=pd.read_csv(RES/'ood_shift_metrics.csv')
cl=pd.read_csv(RES/'closed_loop_metrics.csv')
oodcl=pd.read_csv(RES/'ood_closed_loop_metrics.csv')
tc=pd.read_csv(RES/'throughput_cost_sensitivity.csv')
timing=pd.read_csv(RES/'inference_timing.csv')
summary=json.load(open(RES/'study_summary.json'))
oodsum=json.load(open(RES/'ood_summary.json'))
refs=pd.read_csv(REF).fillna('')

# Helpers

def set_cell_margins(cell, top=50, start=50, bottom=50, end=50):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None:
            node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def set_columns(section, num=2, space_twips=360):
    sectPr=section._sectPr
    cols=sectPr.xpath('./w:cols')
    if cols:
        cols=cols[0]
    else:
        cols=OxmlElement('w:cols'); sectPr.append(cols)
    cols.set(qn('w:num'),str(num)); cols.set(qn('w:space'),str(space_twips)); cols.set(qn('w:equalWidth'),'1')

def set_keep_with_next(p, value=True):
    pPr=p._p.get_or_add_pPr(); el=OxmlElement('w:keepNext'); el.set(qn('w:val'),'1' if value else '0'); pPr.append(el)

def set_keep_lines(p):
    pPr=p._p.get_or_add_pPr(); el=OxmlElement('w:keepLines'); pPr.append(el)

def set_font(run, size=None, bold=None, italic=None, name='Times New Roman'):
    run.font.name=name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic

def shade_cell(cell, fill='D9EAF7'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def remove_table_borders(table):
    tblPr=table._tbl.tblPr; borders=tblPr.first_child_found_in('w:tblBorders')
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for e in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+e; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'nil')

def add_rule_table(doc, headers, rows, widths=None, font_size=7.4, caption=None):
    if caption:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_keep_with_next(p)
        r=p.add_run(caption); set_font(r,8,bold=True)
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    # Word sometimes ignores cell.width inside newspaper columns. Force a fixed table grid.
    if widths:
        twips=[int(w*1440) for w in widths]
        tblPr=t._tbl.tblPr
        layout=tblPr.find(qn('w:tblLayout'))
        if layout is None: layout=OxmlElement('w:tblLayout'); tblPr.append(layout)
        layout.set(qn('w:type'),'fixed')
        tblW=tblPr.find(qn('w:tblW'))
        if tblW is None: tblW=OxmlElement('w:tblW'); tblPr.append(tblW)
        tblW.set(qn('w:w'),str(sum(twips))); tblW.set(qn('w:type'),'dxa')
        grid=t._tbl.tblGrid
        for child in list(grid): grid.remove(child)
        for w in twips:
            gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(w)); grid.append(gc)
        for row in t.rows:
            for j,c in enumerate(row.cells):
                c.width=Inches(widths[j])
                tcPr=c._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
                if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
                tcW.set(qn('w:w'),str(twips[j])); tcW.set(qn('w:type'),'dxa')
    hdr=t.rows[0]; set_repeat_table_header(hdr)
    for j,h in enumerate(headers):
        c=hdr.cells[j]; c.text=''; shade_cell(c,'E6E6E6'); set_cell_margins(c,35,35,35,35)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(str(h)); set_font(r,font_size,bold=True)
    for rr in rows:
        row=t.add_row(); row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        for j,val in enumerate(rr):
            c=row.cells[j]; set_cell_margins(c,30,35,30,35); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                c.width=Inches(widths[j]); tcPr=c._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
                if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
                tcW.set(qn('w:w'),str(int(widths[j]*1440))); tcW.set(qn('w:type'),'dxa')
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j>0 else WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(str(val)); set_font(r,font_size)
    return t

def add_caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(text); set_font(r,7.8)
    return p

def add_figure(doc, filename, caption, width=3.2):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(0)
    p.add_run().add_picture(str(FIG/filename), width=Inches(width))
    add_caption(doc,caption)

def add_heading(doc,text,level=1):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before=Pt(4 if level==1 else 3); p.paragraph_format.space_after=Pt(2)
    set_keep_with_next(p)
    if level==1:
        r=p.add_run(text.upper()); set_font(r,9.4,bold=True)
    else:
        r=p.add_run(text); set_font(r,9.2,bold=True,italic=(level>=3))
    return p

def add_body(doc,text, first_indent=True):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(1.3)
    if first_indent: p.paragraph_format.first_line_indent=Inches(0.14)
    r=p.add_run(text); set_font(r,9.35)
    return p

def add_eq(doc, eq, num):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(eq+f'    ({num})'); set_font(r,9.2,name='Cambria Math')
    return p

def add_bullet(doc,text):
    p=doc.add_paragraph(); p.style=doc.styles['List Bullet']; p.paragraph_format.left_indent=Inches(0.16); p.paragraph_format.first_line_indent=Inches(-0.08); p.paragraph_format.space_after=Pt(1)
    r=p.add_run(text); set_font(r,8.9)
    return p

def add_column_break(doc):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
    p.add_run().add_break(WD_BREAK.COLUMN)
    return p

def f3(x): return f'{float(x):.3f}'
def f4(x): return f'{float(x):.4f}'

# Metrics lookup
lgb=mb[mb.model=='LightGBM'].iloc[0]
cat=mb[mb.model=='CatBoost'].iloc[0]
pgrow=pg[pg.model.str.startswith('PG-CARE')].iloc[0]
phy=pg[pg.model=='LightGBM + physics'].iloc[0]
rf=cm[cm.model=='Random Forest'].iloc[0]
xgb=cm[cm.model=='XGBoost'].iloc[0]
tcmoe=cm[cm.model=='TC-MoE alone'].iloc[0]
pers=cm[cm.model=='Persistence'].iloc[0]
ood_lgb=ood[ood.model=='LightGBM'].iloc[0]; ood_pg=ood[ood.model=='PG-CARE (LightGBM)'].iloc[0]
rawcl=cl[cl.strategy=='XGBoost'].iloc[0]; pgcl=cl[cl.strategy=='PG-CARE risk-calibrated'].iloc[0]; oracle=cl[cl.strategy=='Oracle'].iloc[0]
raw_o=oodcl[oodcl.strategy=='XGBoost raw'].iloc[0]; pg_o=oodcl[oodcl.strategy=='PG-CARE risk-calibrated'].iloc[0]

# Create document
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.48); sec.bottom_margin=Inches(0.50); sec.left_margin=Inches(0.56); sec.right_margin=Inches(0.56)
sec.page_width=Inches(8.5); sec.page_height=Inches(11)

# Styles
styles=doc.styles
normal=styles['Normal']; normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); normal.font.size=Pt(9.35)
normal.paragraph_format.space_after=Pt(1.3); normal.paragraph_format.line_spacing=1.0
for sname in ['List Bullet','List Number']:
    st=styles[sname]; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); st.font.size=Pt(8.9)

# Header
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=header.add_run('MANUSCRIPT PREPARED IN IEEE TRANSACTIONS STYLE'); set_font(r,7.5)
# Footer page field
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run(); fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '; fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end'); run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2); set_font(run,8)

# Title block
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(5)
r=p.add_run('Physics-Guided Conformal Reliability for One-Hour-Ahead Electric-Vehicle Flexibility Forecasting and Risk-Aware V2G Dispatch'); set_font(r,17,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
r=p.add_run('Mahmoud Mohammadnezhad Kiasari and Hamed H. Aly'); set_font(r,10.5)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4)
r=p.add_run('Department of Electrical and Computer Engineering, Dalhousie University, Halifax, NS, Canada'); set_font(r,8.8,italic=True)

abstract=(
'Abstract—Reliable electric-vehicle (EV) flexibility forecasting is difficult because the operational target is strongly constrained by plug-in availability, owner participation, state of charge (SOC), departure requirements, and time-varying grid conditions. This paper develops PG-CARE, a Physics-Guided Conformal Adaptive Reliability Engine that couples modern data-driven forecasting with feasibility projection, calibrated uncertainty sets, and risk-aware vehicle-to-grid (V2G) commitment. A source-informed Python supervisory digital twin extends a previously published MATLAB/Simulink four-EV controller to 100 heterogeneous EVs over 365 days at 15-min resolution, producing 3,504,000 EV-time states without minority-class oversampling. Using a strictly chronological 60/20/20 protocol, the selected LightGBM expert achieves a macro-F1 of 0.7501 and V2G F1 of 0.5471. Physics-feasibility projection raises these metrics to 0.8424 and 0.7969, respectively, while the complete PG-CARE pipeline reaches 0.8482 macro-F1 and 0.7663 Matthews correlation coefficient with 90.12% conformal marginal coverage. A frozen 90-day shifted-regime test yields 0.7970 macro-F1 and 0.7066 V2G F1 without retraining. In closed-loop reserve commitment, a raw flexibility forecast produces negative value under asymmetric shortfall penalties, whereas validation-calibrated PG-CARE converts the same forecasting stack to positive net value in both in-domain and shifted regimes. Results show that physically feasible uncertainty management and decision-aware commitment can be more important than increasing model complexity for trustworthy EV grid services.'
)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.right_indent=Inches(.18); p.paragraph_format.space_after=Pt(3)
r=p.add_run(abstract); set_font(r,8.5)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.right_indent=Inches(.18); p.paragraph_format.space_after=Pt(4)
r=p.add_run('Index Terms—Electric vehicles, vehicle-to-grid, flexibility forecasting, conformal prediction, physics-guided machine learning, digital twin, LightGBM, risk-aware dispatch.'); set_font(r,8.5,italic=True)

# Full-width framework
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
p.add_run().add_picture(str(FIG/'fig1_framework.png'), width=Inches(6.95))
add_caption(doc,'Fig. 1. PG-CARE workflow. A source-informed stochastic EV-fleet digital twin produces chronological operating data; expert forecasts are projected onto the physical feasibility set, wrapped by conformal reliability sets, and passed to a validation-calibrated risk-aware commitment layer.')

# Two columns section
sec2=doc.add_section(WD_SECTION.CONTINUOUS); sec2.top_margin=Inches(.48); sec2.bottom_margin=Inches(.50); sec2.left_margin=Inches(.56); sec2.right_margin=Inches(.56); set_columns(sec2,2,360)

add_heading(doc,'I. Introduction',1)
add_body(doc,'Bidirectional charging converts an EV fleet from a passive load into a time-coupled distributed energy resource. The practical difficulty is not merely deciding whether an EV can charge or discharge at the present instant, but determining whether flexibility will remain physically and behaviorally available when the grid needs it. Recent IEEE Transactions work has therefore moved beyond isolated charging classifiers toward scalable charging networks, dynamic user behavior, aggregate flexibility, stochastic control, battery-aware scheduling, and value-oriented forecasting [2]–[10], [12]–[25].')
add_body(doc,'The starting point of this study is the authors’ previously published MATLAB/Simulink fleet controller [1]. That work introduced owner-selectable participation, minimum-SOC protection, load-dependent V2G/G2V/Idle decisions, and a bidirectional converter-oriented control structure. It was intentionally small: four EVs and four household loads. That scale is useful for controller development, but it is insufficient for evaluating generalization of one-hour-ahead forecasting or the economic consequence of forecast uncertainty.')
add_body(doc,'A second problem arises when forecasting is assessed independently from the decision it supports. A classifier can exhibit good aggregate accuracy while systematically missing the minority V2G state, and a regressor with low root-mean-square error can still create poor reserve commitments when overprediction is penalized more strongly than under-commitment. Modern work on value-oriented forecasting and reserve provision makes this distinction explicit [18], [20], [21]. Thus, a Transactions-level EV forecasting framework should simultaneously respect physical feasibility, quantify uncertainty, remain useful under temporal distribution shift, and demonstrate downstream grid-service value.')
add_body(doc,'This paper addresses that gap with PG-CARE (Physics-Guided Conformal Adaptive Reliability Engine). The contribution is deliberately not a claim that a deeper neural network must outperform tree ensembles. Instead, an expert pool is evaluated chronologically, and the selected predictor is surrounded by reliability mechanisms that encode feasibility and uncertainty before the forecast reaches a dispatch decision. The resulting architecture is intended to fail safely when the data-driven model is unsure or when its preferred state is physically impossible.')
add_body(doc,'The main contributions are fourfold:',first_indent=False)
add_bullet(doc,'A scalable, auditable Python supervisory digital twin extends the logic of the published MATLAB/Simulink controller [1] to 100 heterogeneous EVs over 365 days, generating 3.504 million EV-time states with naturally occurring Idle, G2V, and V2G labels and no SMOTE-based label synthesis.')
add_bullet(doc,'PG-CARE combines validation-based expert selection with a physics-feasibility probability projection and split-conformal prediction sets. The feasibility layer enforces connectivity, participation, SOC, and departure-related constraints before state selection.')
add_bullet(doc,'The study uses chronological evaluation, 1,000 paired bootstrap resamples, and a frozen 90-day shifted-regime stress test with changed mobility, load, PV, and price conditions. No model or threshold is refit on the shifted regime.')
add_bullet(doc,'A closed-loop reserve-commitment experiment evaluates forecast value under asymmetric delivery and shortfall economics. This exposes a result missed by forecast error alone: conservative reliability calibration can increase decision value even while reducing committed energy.')

add_heading(doc,'II. Related Work and Research Gap',1)
add_heading(doc,'A. EV Behavior, Charging, and Flexibility',2)
add_body(doc,'Data-driven EV studies have progressively incorporated richer behavioral structure. Jahangir et al. modeled plug-in behavior using clustered deep learning [2], while Adaptive Charging Networks and ACN-Sim established scalable, data-driven smart-charging and simulation frameworks [3], [4]. Long et al. addressed efficient real-time charging scheduling [5], and Yan et al. incorporated dynamic user behavior using deep reinforcement learning [6]. User anxiety and participation behavior have likewise been embedded explicitly in distributed charging decisions [7]. These studies establish that arrival, departure, energy demand, and owner behavior are first-class state variables rather than nuisance randomness.')
add_heading(doc,'B. Grid-Facing V2G Control',2)
add_body(doc,'The grid-facing literature emphasizes closed-loop performance. Distributed MPC has been used to exploit EV reactive power for voltage regulation [8]; uncertainty has been treated jointly over communication and physical networks [9]; and EV fleets have been embedded in resilience and restoration services [10]. More recent work incorporates bandwidth constraints, cyber/customer interruptions, battery aging, range anxiety, and real-time coordination [13]–[16], [23], [24]. The common implication is that predictive performance should be judged by whether feasible grid support is delivered without violating user or battery constraints.')
add_heading(doc,'C. Modern Forecasting and Generalization',2)
add_body(doc,'Current IEEE Transactions forecasting research increasingly combines multiple temporal scales, graph structure, expert models, robust uncertainty handling, causal/generalization mechanisms, and value-oriented loss functions [11], [19]–[22], [25]. In EV applications, aggregate flexibility is emerging as a more actionable target than state classification alone [12], [17], [18]. However, there remains a methodological gap between (i) highly detailed controller studies with small deterministic fleets and (ii) data-driven forecasts that may violate EV feasibility or become overconfident under regime shift. PG-CARE targets that interface: it uses the published controller logic as the physical prior, but subjects modern forecasters to feasibility projection, conformal uncertainty, out-of-distribution testing, and downstream dispatch evaluation.')
add_heading(doc,'D. Design Requirements Synthesized From 2020–2026 Literature',2)
add_body(doc,'Four design requirements follow from the recent literature. First, EV availability must be stochastic and user-conditioned rather than a fixed daily schedule [2], [6], [7]. Second, forecasting should expose aggregate flexibility and uncertainty because grid services are constrained by the feasible population boundary [12], [17], [18]. Third, model evaluation must include generalization to changed data regimes instead of relying only on interpolation performance [22], [25]. Fourth, forecast quality must be connected to an operational objective; value-oriented methods demonstrate that the ranking of forecasters can change after the downstream decision and its asymmetric costs are considered [20], [21].')
add_column_break(doc)
add_rule_table(doc,['Design need','Representative refs.','PG-CARE response'],[
('User/behavior variability','[2],[6],[7]','Stochastic arrival, departure, trip, policy'),
('Fleet flexibility','[12],[17],[18]','State + SOC + up/down flexibility'),
('Generalization','[22],[25]','Frozen 90-day shifted-regime test'),
('Operational value','[18],[20],[21]','Risk-calibrated reserve commitment')
],widths=[.95,.78,1.34],font_size=6.3,caption='TABLE I  DESIGN REQUIREMENTS DRAWN FROM RECENT IEEE TRANSACTIONS LITERATURE')

add_heading(doc,'III. Source-Informed Python EV Fleet Digital Twin',1)
add_heading(doc,'A. Relationship to the Published MATLAB/Simulink Controller',2)
add_body(doc,'The physical provenance is the controller reported in [1], which used battery SOC, owner participation, minimum SOC, time, and demand thresholds to choose G2V, V2G, or Idle operation. The original implementation also included switching-level converter, filter, and PID details. For year-long fleet learning, reproducing a 10-kHz switching waveform for every vehicle is unnecessary and computationally prohibitive. Accordingly, the present work reconstructs only the supervisory layer at 15-min resolution. It does not claim switching-level or electromagnetic-transient equivalence to the Simulink model.')
add_body(doc,'For EV i, the supervisory battery balance is')
add_eq(doc,'SOCᵢ,t₊₁ = clip[SOCᵢ,t + (ηc Pᶜʰᵍᵢ,t − Pᵈⁱˢᵢ,t/ηd) Δt / Eᵢ]',1)
add_body(doc,'where Eᵢ is usable battery capacity, ηc and ηd are charging/discharging efficiencies, and Δt=0.25 h. Battery capacities are sampled between 40 and 120 kWh and charger ratings from {7.2, 9.6, 11} kW. Each EV is assigned minimum and target SOC limits, charge/discharge efficiencies, and one of three owner policies: disconnected, charge-only, or bidirectional participation.')
add_heading(doc,'B. Stochastic Operating Environment',2)
add_body(doc,'The digital twin contains 100 EVs and 365 days (35,040 fleet time steps). Day-specific arrival time, departure time, trip energy, and required departure SOC are sampled deterministically from seeded stochastic distributions. Native feeder demand includes morning/evening peaks, weekend effects, seasonal modulation, autocorrelated disturbance, and occasional peak events. A stylized PV term and time-of-use price are simulated concurrently. The controller recomputes the first and third quartiles of the previous day’s native-load profile and uses them as low/high demand thresholds, retaining the adaptive threshold concept of [1].')
add_body(doc,'The state yᵢ,t is')
add_eq(doc,'yᵢ,t ∈ {Idle, G2V, V2G}.',2)
add_body(doc,'Charging is activated when an available EV requires energy and either departure urgency, low load, or minimum-SOC proximity makes charging necessary. V2G is permitted only for connected bidirectional participants with adequate SOC margin and no imminent charging urgency, and is activated under high native-load conditions. All infeasible or inactive cases remain Idle. This construction produces natural class imbalance rather than synthesizing minority labels after simulation.')
add_heading(doc,'C. Departure Urgency and Flexibility Labels',2)
add_body(doc,'The supervisory logic estimates the average charging power required to reach the driver-specific departure target. For remaining time Tᵣ and energy deficit Ereq=max[0,(SOCreq−SOC)E], the urgency variable is u=Ereq/max(Tᵣ,Δt). Charging is flagged as urgent when u exceeds 72% of the rated charger power. This prevents V2G decisions that would consume energy needed for an imminent departure.')
add_eq(doc,'uᵢ,t = max[0,(SOCreqᵢ,t−SOCᵢ,t)Eᵢ] / max(Tᵣᵢ,t,Δt).',3)
add_body(doc,'The auxiliary flexibility labels are generated from the same feasibility state rather than from realized future dispatch power. At t+H, upward flexibility is the charger rating only if the EV is connected, bidirectionally enrolled, and has a 0.06-p.u. SOC margin above its owner minimum. Downward flexibility is the charger rating when a connected vehicle remains below its required departure SOC by more than 0.01 p.u.:')
add_eq(doc,'P↑ᵢ,t+H = cᵢ·1{policy=V2G}·1{SOC> SOCmin+0.06}·Pratedᵢ,',4)
add_eq(doc,'P↓ᵢ,t+H = cᵢ·1{SOC< SOCreq−0.01}·Pratedᵢ.',5)
add_body(doc,'These targets distinguish “what state will the controller select?” from “how much physically feasible charging/discharging capability exists?” This distinction is central to the later closed-loop experiment because reserve commitment uses aggregate P↑ rather than the categorical state alone.')
add_rule_table(doc,['Quantity','Value'],[
('Fleet size','100 EVs'),('Duration / resolution','365 days / 15 min'),('EV-time states','3,504,000'),('Idle / G2V / V2G','84.81 / 11.38 / 3.81%'),('Battery capacity','40–120 kWh'),('Charger ratings','7.2, 9.6, 11 kW'),('Owner policies','disconnect / charge-only / bidirectional'),('Annual G2V energy','855.26 MWh'),('Annual V2G energy','289.04 MWh')
],widths=[1.55,1.55],font_size=7.2,caption='TABLE II  DIGITAL-TWIN CONFIGURATION AND ANNUAL OPERATING STATISTICS')
add_figure(doc,'fig2_week_operation.png','Fig. 2. Representative one-week operation of the stochastic fleet, showing native demand, aggregate EV power, feeder power, connected vehicles, and mean SOC.',3.18)
add_figure(doc,'fig3_class_distribution.png','Fig. 3. Natural operating-state distribution over the full one-year digital twin. The minority V2G state is retained without oversampling.',3.18)

add_heading(doc,'IV. PG-CARE Forecasting and Reliability Architecture',1)
add_heading(doc,'A. Forecasting Tasks and Temporal Features',2)
add_body(doc,'At decision time t, each learning sample contains the previous L=8 supervisory intervals (2 h) and predicts conditions H=4 intervals (1 h) ahead. Sixteen features per interval represent SOC, connectivity, one-hot current state, normalized time-to-departure, required SOC, normalized native load and price, cyclic time-of-day, load position relative to the Q1/Q3 range, battery capacity, charger power, minimum SOC, and owner policy. The principal classification map is')
add_eq(doc,'fθ : Xᵢ,t−L+1:t → ŷᵢ,t+H.',6)
add_body(doc,'Two auxiliary regressions estimate future SOC and feasible upward/downward flexibility. Aggregate upward flexibility is used by the reserve-commitment experiment, so the predictive layer is evaluated both statistically and through its downstream operating value.')
add_heading(doc,'B. Expert Pool and Validation-Based Selection',2)
add_body(doc,'The expert pool contains logistic regression, random forest, XGBoost, LightGBM, CatBoost, and a temporal convolutional mixture-of-experts (TC-MoE). The TC-MoE uses two one-dimensional convolutional layers, three learned regime experts, and a gating network; its multi-task head jointly estimates state probabilities, SOC, and normalized flexibility. Importantly, model complexity is not forced into the final architecture. On the validation period, LightGBM is selected as the classification expert, while XGBoost is retained for the continuous flexibility regressions. TC-MoE remains in the paper as a negative result because it does not outperform the tree ensembles on this structured supervisory dataset.')
add_rule_table(doc,['Expert','Configuration'],[
('Random forest','35 trees, depth 12, balanced subsampling'),
('XGBoost classifier','65 trees, depth 5, η=0.10, class weighting'),
('LightGBM','160 trees, 31 leaves, η=0.05, class weighting'),
('CatBoost','180 iterations, depth 7, η=0.07'),
('TC-MoE','2×Conv1D(32), 3 experts, learned gate')
],widths=[.92,2.13],font_size=6.4,caption='TABLE III  REPRESENTATIVE EXPERT CONFIGURATIONS')
add_heading(doc,'C. Physics-Feasibility Projection',2)
add_body(doc,'Let p=[pIdle,pG2V,pV2G] denote the classifier probabilities and m∈{0,1}³ a feasibility mask computed from future connectivity, owner policy, SOC margin, charger limit, and departure requirement. Infeasible class probabilities are projected out before a state is selected:')
add_eq(doc,'p̃ = (m ⊙ p) / [1ᵀ(m ⊙ p)].',7)
add_body(doc,'The mask prevents V2G for disconnected or non-participating EVs and suppresses discharge when the available energy cannot preserve the prescribed SOC margin. If every non-Idle action is infeasible, the distribution collapses to Idle. This differs from adding a soft penalty during training: feasibility is enforced at inference regardless of classifier confidence.')
add_heading(doc,'D. Split-Conformal Reliability Sets',2)
add_body(doc,'A calibration threshold is derived only from the chronological validation period using the nonconformity score s=1−p̃y for the true class. The higher 0.90 empirical quantile q0.90 is computed on validation samples, and τconf=1−q0.90. For each test sample, PG-CARE forms a set of labels whose projected probabilities exceed this frozen threshold. The implementation uses')
add_eq(doc,'Γ(x) = {k : p̃k(x) ≥ τconf}.',8)
add_body(doc,'The validation-selected LightGBM threshold is τconf=0.6287. Singleton sets are accepted. Non-singleton or empty sets trigger a conservative fallback: an urgent connected vehicle is assigned G2V; otherwise the controller chooses Idle. On the held-out chronological test period, the resulting marginal coverage is 0.9012 and 4.56% of samples invoke the ambiguous-set fallback.')
add_heading(doc,'E. Safety Semantics of the Fallback',2)
add_body(doc,'The fallback is intentionally asymmetric. An ambiguous forecast does not authorize V2G, because a false discharge commitment can violate owner SOC or create reserve shortfall. G2V is selected only when the departure-urgency rule is already active; otherwise PG-CARE defaults to Idle. Thus, uncertainty is converted into a controllable loss of opportunistic grid service rather than an infeasible battery action. This design is one reason balanced accuracy can decline slightly after conformal fallback even while macro-F1, MCC, and realized decision value improve.')
add_heading(doc,'F. Risk-Aware Reserve Commitment',2)
add_body(doc,'Forecasting available upward flexibility is not equivalent to deciding how much reserve should be committed. Let F̂t be the aggregate predicted flexibility and ρ∈[0,1] a commitment factor. PG-CARE selects ρ only on validation data by maximizing a value function with delivered-reserve revenue r and shortfall penalty c:')
add_eq(doc,'ρ* = arg maxρ Σt [ r·min(ρF̂t,Ft) − c·max(0,ρF̂t−Ft) ].',9)
add_body(doc,'The validation optimum is ρ*=0.10 for the study assumptions r=$0.18/kWh and c=$0.55/kWh. These values are controlled scenario economics rather than a claim about a specific wholesale market. The same frozen factor is used for the chronological test and the shifted-regime experiment.')

add_heading(doc,'V. Experimental Protocol',1)
add_body(doc,'All reported headline results use a strict chronological split: the first 60% of the 365-day simulation is the training period, the next 20% is validation, and the final 20% is test. Deterministic sample budgets are 30,000, 10,000, and 20,000 EV-window examples, respectively. Normalization parameters, model selection, conformal thresholds, and risk factors are determined without test-period access. No random window split is used for headline claims.')
add_body(doc,'Classification metrics include accuracy, balanced accuracy, macro-F1, Matthews correlation coefficient (MCC), and per-class F1. The primary metrics are macro-F1 and MCC because 84.81% of EV-time states are Idle. Continuous SOC and flexibility models are evaluated with mean absolute error. For statistical uncertainty, 1,000 paired bootstrap resamples of the test predictions quantify confidence intervals and PG-CARE-minus-LightGBM deltas.')
add_body(doc,'Generalization is tested by freezing the complete in-domain pipeline and applying it without retraining to a separately simulated 90-day regime with later arrivals, earlier departures, longer trips, heavier and peakier native demand, lower PV, and more volatile prices. This is intentionally an out-of-distribution synthetic stress test, not external field validation.')
add_rule_table(doc,['Shifted variable','In-domain generator','OOD generator'],[
('Weekday arrival mean','17.4 h','18.8 h'),('Weekday departure mean','7.4 h','6.7 h'),('Weekday trip log-mean','14 kWh','18 kWh'),('PV scale','1.20×N','0.95×N'),('Evening load coefficient','0.95','1.08'),('Peak-price adder','$0.13/kWh','$0.17/kWh')
],widths=[1.10,.98,.98],font_size=6.2,caption='TABLE IV  SELECTED SHIFTED-REGIME GENERATOR CHANGES')
add_body(doc,'The OOD experiment reuses the static EV capacities, charger ratings, minimum SOC values, and owner policies from the in-domain fleet while changing daily mobility and exogenous operating conditions. In-domain normalization statistics, trained classifier/regressors, feasibility rules, conformal thresholds, and the reserve commitment factor are all frozen. Consequently, the stress test measures robustness to a changed temporal regime rather than benefit from adaptation to the test distribution.')
add_body(doc,'The software environment uses Python 3.13.5, NumPy 2.3.5, pandas 2.2.3, scikit-learn 1.8.0, XGBoost 3.1.3, PyTorch 2.10.0+cpu, together with LightGBM and CatBoost. Random seeds, split indices, trained models, prediction audits, generated data, and figure scripts are retained in the accompanying reproducibility bundle.')

add_heading(doc,'VI. Results',1)
add_heading(doc,'A. Chronological Operating-State Forecasting',2)
rows=[]
for name,row in [('Persistence',pers),('Random Forest',rf),('XGBoost',xgb),('TC-MoE',tcmoe),('LightGBM',lgb),('PG-CARE',pgrow)]:
    rows.append((name,f3(row['macro_f1']),f3(row['mcc']),f3(row['v2g_f1'])))
add_rule_table(doc,['Model','Macro-F1','MCC','V2G F1'],rows,widths=[1.20,.62,.62,.62],font_size=6.9,caption='TABLE V  CHRONOLOGICAL ONE-HOUR-AHEAD CLASSIFICATION RESULTS')
add_body(doc,f'LightGBM is the strongest raw expert with macro-F1={lgb.macro_f1:.4f}, MCC={lgb.mcc:.4f}, and V2G F1={lgb.v2g_f1:.4f}. The temporal TC-MoE obtains macro-F1={tcmoe.macro_f1:.4f}; thus, the validation procedure does not assign it the final decision role. After reliability processing, PG-CARE reaches macro-F1={pgrow.macro_f1:.4f}, MCC={pgrow.mcc:.4f}, and V2G F1={pgrow.v2g_f1:.4f}. The absolute gains over raw LightGBM are +{pgrow.macro_f1-lgb.macro_f1:.4f}, +{pgrow.mcc-lgb.mcc:.4f}, and +{pgrow.v2g_f1-lgb.v2g_f1:.4f}, respectively.')
add_figure(doc,'fig4_model_comparison.png','Fig. 4. Chronological one-hour-ahead model comparison. Aggregate accuracy alone is insufficient; macro-F1, MCC, and V2G F1 expose minority-state reliability.',3.18)
add_figure(doc,'fig10_confusion.png','Fig. 5. Row-normalized confusion matrices for the raw selected expert and PG-CARE. The reliability layers primarily reduce infeasible minority-state errors while retaining strong Idle performance.',3.18)

add_heading(doc,'B. Reliability Ablation and Statistical Significance',2)
add_rule_table(doc,['Variant','Macro-F1','MCC','V2G F1'],[
('LightGBM',f3(lgb.macro_f1),f3(lgb.mcc),f3(lgb.v2g_f1)),
('+ Physics',f3(phy.macro_f1),f3(phy.mcc),f3(phy.v2g_f1)),
('+ Conformal (PG-CARE)',f3(pgrow.macro_f1),f3(pgrow.mcc),f3(pgrow.v2g_f1))
],widths=[1.38,.57,.57,.57],font_size=6.9,caption='TABLE VI  ABLATION OF PG-CARE RELIABILITY LAYERS')
add_body(doc,'The largest improvement is produced by the feasibility projection, not by adding a deeper predictor. Physics projection raises V2G F1 from 0.5471 to 0.7969 while also increasing macro-F1 and MCC. Conformal fallback further increases accuracy and macro-F1, although balanced accuracy falls slightly relative to physics-only operation because ambiguous cases are intentionally resolved conservatively toward G2V/Idle.')
# CI text
ci_pg={r.metric:r for _,r in ci[ci.model=='PG-CARE'].iterrows()}
add_body(doc,f'Bootstrap uncertainty confirms that the gains are not attributable to one test draw. PG-CARE macro-F1 is {ci_pg["macro_f1"].estimate:.4f} (95% CI {ci_pg["macro_f1"].ci_low:.4f}–{ci_pg["macro_f1"].ci_high:.4f}), MCC is {ci_pg["mcc"].estimate:.4f} ({ci_pg["mcc"].ci_low:.4f}–{ci_pg["mcc"].ci_high:.4f}), and V2G F1 is {ci_pg["v2g_f1"].estimate:.4f} ({ci_pg["v2g_f1"].ci_low:.4f}–{ci_pg["v2g_f1"].ci_high:.4f}). The paired 1,000-resample deltas versus raw LightGBM are +0.0981 macro-F1 [0.0902, 0.1054], +0.0896 MCC [0.0811, 0.0974], and +0.2464 V2G F1 [0.2255, 0.2664].')
add_body(doc,'The 90.12% marginal test coverage is close to the requested 90% calibration level, while only 4.56% of test windows yield non-singleton or empty reliability sets. On the shifted regime, coverage decreases modestly to 89.42% and the ambiguous-set rate remains 4.20%. This behavior is preferable to interpreting the maximum softmax/tree probability as a calibrated guarantee: PG-CARE has an explicit mechanism for declining a high-risk V2G action when uncertainty is unresolved.')
add_figure(doc,'fig5_reliability_ablation.png','Fig. 6. Reliability-layer ablation. Feasibility projection supplies the dominant V2G reliability gain; conformal fallback improves conservative overall classification.',3.18)

add_heading(doc,'C. SOC and Flexibility Regression',2)
rr=[]
for _,r in reg.iterrows():
    unit='SOC fraction' if r.target=='SOC' else 'kW'
    rr.append((r['model'],r['target'],f'{r.mae:.4f} {unit}'))
add_rule_table(doc,['Model','Target','MAE'],rr,widths=[.78,1.00,1.22],font_size=6.8,caption='TABLE VII  AUXILIARY REGRESSION PERFORMANCE')
add_body(doc,'The continuous targets reinforce the expert-selection result. XGBoost attains SOC MAE 0.0212 (approximately 2.12 percentage points), upward-flexibility MAE 0.953 kW, and downward-flexibility MAE 0.841 kW, whereas TC-MoE errors are materially larger. Consequently, the closed-loop experiment uses the XGBoost flexibility regressor rather than forcing the deep model into the operational pathway.')

add_heading(doc,'D. Shifted-Regime Generalization Without Retraining',2)
add_rule_table(doc,['Model','Macro-F1','MCC','V2G F1'],[
('LightGBM',f3(ood_lgb.macro_f1),f3(ood_lgb.mcc),f3(ood_lgb.v2g_f1)),
('PG-CARE',f3(ood_pg.macro_f1),f3(ood_pg.mcc),f3(ood_pg.v2g_f1))
],widths=[1.35,.58,.58,.58],font_size=7.0,caption='TABLE VIII  FROZEN 90-DAY SHIFTED-REGIME RESULTS')
add_body(doc,f'The shifted regime is intentionally harsher than the training environment. Raw LightGBM falls to macro-F1={ood_lgb.macro_f1:.4f} and V2G F1={ood_lgb.v2g_f1:.4f}. With no retraining, PG-CARE reaches macro-F1={ood_pg.macro_f1:.4f} and V2G F1={ood_pg.v2g_f1:.4f}. The complete pipeline therefore retains {100*ood_pg.macro_f1/pgrow.macro_f1:.1f}% of its in-domain macro-F1 while increasing shifted-regime V2G F1 by {ood_pg.v2g_f1-ood_lgb.v2g_f1:.4f} absolute. Conformal marginal coverage in this shifted experiment is {oodsum["lgbm_conformal_coverage"]:.4f}.')
add_figure(doc,'fig6_ood_generalization.png','Fig. 7. Frozen shifted-regime evaluation. Later arrivals, earlier departures, longer trips, heavier loading, lower PV, and price volatility reduce raw-model reliability; PG-CARE preserves substantially stronger V2G detection.',3.18)

add_column_break(doc)
add_heading(doc,'E. Closed-Loop Reserve Commitment',2)
add_rule_table(doc,['Strategy','Net value ($)','Delivered (MWh)','Shortfall (MWh)'],[
('Persistence',f'{cl[cl.strategy=="Persistence"].iloc[0]["net_value_$"]:.0f}',f'{cl[cl.strategy=="Persistence"].iloc[0].delivered_mwh:.2f}',f'{cl[cl.strategy=="Persistence"].iloc[0].shortfall_mwh:.2f}'),
('Raw XGBoost',f'{rawcl["net_value_$"]:.0f}',f'{rawcl.delivered_mwh:.2f}',f'{rawcl.shortfall_mwh:.2f}'),
('PG-CARE',f'{pgcl["net_value_$"]:.0f}',f'{pgcl.delivered_mwh:.2f}',f'{pgcl.shortfall_mwh:.2f}'),
('Oracle',f'{oracle["net_value_$"]:.0f}',f'{oracle.delivered_mwh:.2f}',f'{oracle.shortfall_mwh:.2f}')
],widths=[.95,.70,.72,.72],font_size=6.5,caption='TABLE IX  IN-DOMAIN CLOSED-LOOP RESERVE-COMMITMENT VALUE')
add_body(doc,f'Raw XGBoost has the lowest flexibility RMSE among learned models ({rawcl.flex_rmse_kw:.1f} kW) yet produces ${rawcl["net_value_$"]:.0f} net value because 11.49 MWh of committed reserve is unavailable at delivery. PG-CARE deliberately commits only a validation-calibrated fraction of predicted flexibility, reducing shortfall to {pgcl.shortfall_mwh:.2f} MWh and yielding +${pgcl["net_value_$"]:.0f}. This is a central result: a lower point-forecast error does not guarantee a better grid-service decision under asymmetric consequences.')
add_body(doc,f'The same frozen policy remains beneficial under the shifted regime: raw XGBoost yields ${raw_o["net_value_$"]:.0f}, whereas PG-CARE yields +${pg_o["net_value_$"]:.0f} with shortfall reduced from {raw_o.shortfall_mwh:.2f} to {pg_o.shortfall_mwh:.2f} MWh. The oracle result bounds the attainable value and is not a deployable benchmark.')
add_figure(doc,'fig7_closed_loop_value.png','Fig. 8. Net closed-loop reserve value under asymmetric shortfall penalties. Risk-aware commitment changes the sign of realized value in both in-domain and shifted-regime tests.',3.18)
add_figure(doc,'fig8_shortfall.png','Fig. 9. Reserve-delivery shortfall. Raw forecasts overcommit available upward flexibility; the frozen PG-CARE risk factor trades volume for reliability.',3.18)

add_heading(doc,'F. Throughput-Cost Sensitivity',2)
# get PG-Care rows
pg_tc=tc[tc.strategy=='PG-CARE risk-calibrated']
vals=[]
for _,r in pg_tc.iterrows(): vals.append((f'${r["throughput_cost_$/kWh"]:.02f}',f'${r["net_after_throughput_$"]:.0f}'))
add_rule_table(doc,['Throughput proxy ($/kWh)','PG-CARE net after proxy'],vals,widths=[1.45,1.45],font_size=6.9,caption='TABLE X  BATTERY-THROUGHPUT COST SENSITIVITY')
add_body(doc,'Because a full electrochemical degradation model is outside the present evidence boundary, battery wear is not converted into a claimed life-cycle cost. Instead, a transparent throughput-cost sensitivity is reported. PG-CARE remains positive at $0.03/kWh (+$369) and $0.06/kWh (+$40), but becomes negative at $0.10/kWh. The implied in-domain break-even throughput proxy is approximately $0.064/kWh. This result should be interpreted as sensitivity, not battery-aging validation.')
add_figure(doc,'fig9_throughput_sensitivity.png','Fig. 10. Net reserve value after an assumed battery-throughput cost. The sensitivity identifies the range in which the conservative grid-service policy remains economically positive under the study assumptions.',3.18)
add_heading(doc,'G. Computational Cost',2)
add_body(doc,f'Training the selected LightGBM classifier requires {lgb.train_s:.2f} s in the recorded CPU environment. Repeated inference timing over the 20,000-sample chronological test set gives a median of {timing[timing.component=="LightGBM classifier"].iloc[0].median_total_ms:.1f} ms in total ({timing[timing.component=="LightGBM classifier"].iloc[0].median_us_per_sample:.2f} μs/sample). The XGBoost upward-flexibility regressor requires {timing[timing.component=="XGBoost flexibility"].iloc[0].median_total_ms:.1f} ms for the same 20,000 feature vectors. These software timings are orders of magnitude below the 15-min supervisory interval, but they are not HIL or embedded-controller latency measurements.')

add_heading(doc,'VII. Discussion',1)
add_heading(doc,'A. Reliability Layers Matter More Than Model Depth',2)
add_body(doc,'The experiments do not support a simplistic “newest deep model wins” narrative. TC-MoE underperforms LightGBM and XGBoost on the structured supervisory data. The stronger contribution is architectural: the raw expert supplies discriminative probabilities, the feasibility layer removes physically impossible actions, conformal sets identify unreliable decisions, and the commitment layer explicitly prices the consequence of overestimating flexibility. This decomposition also makes failure modes easier to audit.')
add_heading(doc,'B. Why Physics Projection Helps the Minority V2G State',2)
add_body(doc,'V2G occupies only 3.81% of EV-time states, yet false positives are especially dangerous because they imply unavailable discharge capability. Many raw classification errors occur in regions where owner policy, connectivity, or SOC already excludes V2G. Projection removes those states from the hypothesis space after learning rather than asking the classifier to relearn deterministic rules from imbalanced samples. The resulting V2G F1 gain of roughly 0.25 over raw LightGBM is therefore consistent with the controller structure rather than an opaque post-processing boost.')
add_heading(doc,'C. Forecast Accuracy Versus Decision Value',2)
add_body(doc,'The closed-loop results align with recent value-oriented forecasting research [18], [20], [21]. XGBoost’s lower flexibility RMSE encourages aggressive commitments but incurs costly shortfalls. PG-CARE’s 0.10 validation-selected commitment factor is conservative by design and produces positive realized value under the study penalty ratio. This does not prove that 0.10 is universally optimal; it demonstrates that the forecasting objective and the downstream risk model cannot be separated when EV flexibility is sold as a grid service.')
add_heading(doc,'D. Generalization and Transactions-Level Evidence',2)
add_body(doc,'The 90-day OOD experiment is important because the model is not adapted after mobility, load, PV, and price statistics shift. The improvement from 0.6840 to 0.7970 macro-F1 and from 0.4007 to 0.7066 V2G F1 indicates that explicit feasibility remains useful when the statistical expert deteriorates. Nevertheless, synthetic OOD evidence is not equivalent to cross-site field validation. A future Transactions extension should pair the present reproducible digital twin with site-specific charging records and distribution-network power-flow or HIL validation.')
add_heading(doc,'E. Relationship to Existing Transactions-Level Methods',2)
add_body(doc,'PG-CARE is complementary to, rather than a replacement for, advanced EV scheduling and forecasting methods. Aggregate-flexibility methods [12], [17], [18] provide richer feasible-set descriptions than the scalar up/down labels used here; graph-based and multi-time-scale predictors [11], [19] can replace the selected tree expert when spatial or station-level structure warrants it; and MPC-based scheduling [8], [24] can replace the simple reserve-commitment layer. The contribution of the present study is the reliability interface connecting any such predictor to physical feasibility, calibrated uncertainty, and decision value. The architecture is therefore modular: the expert and downstream optimizer can be upgraded without changing the evidence discipline around chronological validation and safety projection.')
add_body(doc,'This modularity also explains the deliberate negative result for TC-MoE. On this dataset, static battery/user variables and deterministic feasibility structure carry substantial predictive information, so tree ensembles are difficult to beat with a compact temporal neural network. Hiding that outcome and reporting only the deeper model would weaken reproducibility. Instead, expert selection is treated as an empirical design choice, while novelty is concentrated in the physics/conformal/value layers that remain applicable when the expert changes.')

add_heading(doc,'VIII. Limitations, Reproducibility, and Evidence Boundary',1)
add_body(doc,'First, the reported digital twin is source-informed and stochastic but synthetic; ACN and other public charging datasets are discussed in the literature but were not used to fit the reported results. Second, the Python model is a 15-min supervisory reconstruction of the published controller and does not reproduce 10-kHz switching, converter harmonics, or electromagnetic transients. Third, no distribution power-flow, voltage, transformer thermal, or hardware-in-the-loop result is claimed. Fourth, the reserve prices and penalties are controlled study assumptions, and the battery-wear analysis is a throughput-cost sensitivity rather than an electrochemical aging model.')
add_body(doc,'These boundaries are intentional. The reproducibility package preserves the original Simulink artifact, the Python generator, trained models, deterministic split indices, prediction audits, all CSV results, figure code, package versions, and SHA-256 manifest. The study can therefore be rerun end-to-end using a single orchestration script. This also enables future work to replace the synthetic mobility/load modules or the supervisory plant without changing the forecasting/reliability evaluation pipeline.')

add_heading(doc,'IX. Conclusion',1)
add_body(doc,'This paper developed PG-CARE, a physics-guided conformal reliability framework for one-hour-ahead EV operating-state and flexibility forecasting. A source-informed Python reconstruction extended a published four-EV MATLAB/Simulink supervisory controller to a one-year, 100-EV stochastic fleet with naturally imbalanced Idle/G2V/V2G states. Under chronological testing, a validation-selected LightGBM expert achieved 0.7501 macro-F1 and 0.5471 V2G F1; feasibility projection and conformal safety increased the complete PG-CARE performance to 0.8482 macro-F1, 0.7663 MCC, and 0.7935 V2G F1 with approximately 90% conformal coverage. The same frozen pipeline maintained 0.7970 macro-F1 on a shifted 90-day regime. Most importantly, closed-loop reserve tests showed that lower point-forecast error did not guarantee positive operating value: raw flexibility forecasts overcommitted reserve, while validation-calibrated PG-CARE produced positive net value and much lower shortfall in both regimes. The results support a shift from model-centric EV forecasting toward feasibility-, uncertainty-, and decision-aware forecasting architectures. Future work should validate the same framework using multi-site measured charging records, distribution-network power flow, detailed battery aging, and controller/hardware-in-the-loop experiments.')

# Data availability
add_heading(doc,'Data and Code Availability',1)
add_body(doc,'The complete reproducibility package accompanying this manuscript contains the source-informed digital-twin generator, model-training and OOD scripts, frozen trained models, generated supervisory datasets, prediction audits, all reported result tables, figures, the original Simulink source artifact used for provenance, environment information, and a SHA-256 manifest. Because the simulation is deterministic under the recorded seeds and split indices, the principal tables and figures can be regenerated from the package.',first_indent=False)

# References
add_heading(doc,'References',1)
for _,r in refs.iterrows():
    no=int(r['no']); authors=str(r['authors']); title=str(r['title']); journal=str(r['journal']); vol=str(r['volume']); issue=str(r['issue']); pages=str(r['pages']); year=int(r['year']); doi=str(r['doi'])
    # IEEE-ish formatting, all refs 2020-2026 by design.
    parts=[f'[{no}] {authors}, “{title},” {journal}']
    if vol: parts.append(f', vol. {vol}')
    if issue:
        try: issue=str(int(float(issue)))
        except: pass
        parts.append(f', no. {issue}')
    if pages:
        if pages.isdigit() and len(pages)>=6 and '-' not in pages: parts.append(f', Art. no. {pages}')
        else: parts.append(f', pp. {pages}')
    parts.append(f', {year}')
    if doi: parts.append(f', doi: {doi}')
    parts.append('.')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.18); p.paragraph_format.space_after=Pt(.6); p.paragraph_format.line_spacing=0.95
    rr=p.add_run(''.join(parts)); set_font(rr,7.5)

# Keep tables within margins, tweak table borders
for table in doc.tables:
    table.autofit=False
    # top and bottom rule + header bottom; light grid for legibility
    tblPr=table._tbl.tblPr
    borders=tblPr.first_child_found_in('w:tblBorders')
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge,val,size in [('top','single','8'),('bottom','single','8'),('insideH','single','2'),('insideV','nil','0'),('left','nil','0'),('right','nil','0')]:
        el=borders.find(qn('w:'+edge))
        if el is None: el=OxmlElement('w:'+edge); borders.append(el)
        el.set(qn('w:val'),val); el.set(qn('w:sz'),size); el.set(qn('w:color'),'808080')

# Document properties
doc.core_properties.title='Physics-Guided Conformal Reliability for One-Hour-Ahead Electric-Vehicle Flexibility Forecasting and Risk-Aware V2G Dispatch'
doc.core_properties.subject='IEEE Transactions-style EV flexibility forecasting manuscript'
doc.core_properties.author='Mahmoud Mohammadnezhad Kiasari; Hamed H. Aly'
doc.core_properties.keywords='EV, V2G, conformal prediction, physics-guided ML, flexibility forecasting, digital twin'

doc.save(OUT)
print(OUT)
